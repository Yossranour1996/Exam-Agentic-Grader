"""Small, consistent file I/O helpers."""

import json
from base64 import b64encode
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook, load_workbook

from src.utils.logging import Logger


def _error(message: str, error: Exception, logger: Logger | None) -> None:
    text = f"{message}: {error}"
    print(text)
    if logger:
        logger.log(text, level="error")


def read_text(filepath: str | Path, logger: Logger | None = None, encoding="utf-8") -> str | None:
    try:
        return Path(filepath).read_text(encoding=encoding)
    except Exception as error:
        _error(f"Error reading {filepath}", error, logger)
        return None


def write_text(filepath: str | Path, content: str, logger: Logger | None = None, encoding="utf-8", append=False) -> bool:
    path = Path(filepath)
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("a" if append else "w", encoding=encoding) as file:
            file.write(content)
        return True
    except Exception as error:
        _error(f"Error writing {filepath}", error, logger)
        return False


def read_json(filepath: str | Path, logger: Logger | None = None) -> Any | None:
    try:
        text = Path(filepath).read_text(encoding="utf-8").strip()
        return json.loads(text) if text else None
    except FileNotFoundError:
        return None
    except Exception as error:
        _error(f"Error parsing JSON from {filepath}", error, logger)
        return None


def read_yaml(filepath: str | Path, logger: Logger | None = None) -> Any | None:
    try:
        return yaml.safe_load(Path(filepath).read_text(encoding="utf-8"))
    except Exception as error:
        _error(f"Error parsing YAML from {filepath}", error, logger)
        return None


def write_json(filepath: str | Path, data: Any, logger: Logger | None = None, indent=2, ensure_ascii=False) -> bool:
    path = Path(filepath)
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        final = data
        existing = read_json(path, logger) if path.exists() else None
        if isinstance(existing, list):
            final = existing + (data if isinstance(data, list) else [data])
        path.write_text(json.dumps(final, ensure_ascii=ensure_ascii, indent=indent), encoding="utf-8")
        return True
    except Exception as error:
        _error(f"Error writing JSON to {filepath}", error, logger)
        return False


def read_excel(filepath: str | Path, logger: Logger | None = None) -> dict[str, list[dict[str, Any]]] | None:
    try:
        workbook = load_workbook(filepath, data_only=True)
        result = {}
        for name in workbook.sheetnames:
            rows = list(workbook[name].iter_rows(values_only=True))
            headers = [str(value) if value is not None else f"Column_{i}" for i, value in enumerate(rows[0])] if rows else []
            result[name] = [dict(zip(headers, row)) for row in rows[1:]]
        return result
    except Exception as error:
        _error(f"Error reading Excel file {filepath}", error, logger)
        return None


def write_excel(filepath: str | Path, data: dict[str, list[dict[str, Any]]], logger: Logger | None = None) -> bool:
    path = Path(filepath)
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        existing = read_excel(path, logger) if path.exists() else {}
        combined = {name: (existing or {}).get(name, []) + rows for name, rows in data.items()}
        for name, rows in (existing or {}).items():
            combined.setdefault(name, rows)
        workbook = Workbook()
        workbook.remove(workbook.active)
        for name, rows in combined.items():
            sheet = workbook.create_sheet(name)
            if rows:
                headers = list(dict.fromkeys(key for row in rows for key in row))
                sheet.append(headers)
                for row in rows:
                    sheet.append([row.get(header, "") for header in headers])
        workbook.save(path)
        return True
    except Exception as error:
        _error(f"Error writing Excel file {filepath}", error, logger)
        return False


def read_base64(filepath: str | Path, logger: Logger | None = None) -> str | None:
    try:
        return b64encode(Path(filepath).read_bytes()).decode("ascii")
    except Exception as error:
        _error(f"Error reading {filepath} as base64", error, logger)
        return None


def file_exists(filepath: str | Path) -> bool:
    return Path(filepath).exists()


def delete_file(filepath: str | Path, logger: Logger | None = None) -> bool:
    try:
        Path(filepath).unlink(missing_ok=True)
        return True
    except Exception as error:
        _error(f"Error deleting {filepath}", error, logger)
        return False
